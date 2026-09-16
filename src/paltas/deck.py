"""Utilidades para construir presentaciones .pptx.

Las presentaciones se generan desde los JSON de metricas y las figuras del
repositorio, no se editan a mano. Asi no pueden quedar desincronizadas con los
resultados: si cambia un entrenamiento, se regenera el deck y las cifras se
actualizan solas.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# Paleta del proyecto
VERDE = RGBColor(0x2E, 0x6B, 0x2F)
VERDE_CLARO = RGBColor(0x8F, 0xB3, 0x39)
GRIS = RGBColor(0x44, 0x4A, 0x52)
GRIS_CLARO = RGBColor(0x8A, 0x90, 0x98)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
AZUL = RGBColor(0x1F, 0x6F, 0xB2)
ROJO = RGBColor(0xD1, 0x49, 0x5B)
MORADO = RGBColor(0x6A, 0x4C, 0x93)

ANCHO = Inches(13.333)
ALTO = Inches(7.5)


def nueva_presentacion() -> Presentation:
    prs = Presentation()
    prs.slide_width = ANCHO
    prs.slide_height = ALTO
    return prs


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _caja(slide, x, y, w, h, texto, size=18, bold=False, color=GRIS,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lineas = texto.split("\n")
    for i, linea in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = linea
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb


def slide_portada(prs, titulo, subtitulo, autor, fecha, pie=""):
    s = _blank(prs)
    fondo = s.shapes.add_shape(1, 0, 0, ANCHO, ALTO)
    fondo.fill.solid()
    fondo.fill.fore_color.rgb = VERDE
    fondo.line.fill.background()
    fondo.shadow.inherit = False

    barra = s.shapes.add_shape(1, 0, Inches(4.55), ANCHO, Inches(0.06))
    barra.fill.solid()
    barra.fill.fore_color.rgb = VERDE_CLARO
    barra.line.fill.background()
    barra.shadow.inherit = False

    _caja(s, Inches(1.0), Inches(2.5), Inches(11.3), Inches(1.9), titulo,
          size=44, bold=True, color=BLANCO)
    _caja(s, Inches(1.0), Inches(4.85), Inches(11.3), Inches(1.0), subtitulo,
          size=20, color=RGBColor(0xDA, 0xE8, 0xC8))
    _caja(s, Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.8),
          f"{autor}  ·  {fecha}" + (f"\n{pie}" if pie else ""),
          size=14, color=RGBColor(0xC2, 0xD6, 0xAE))
    return s


def slide_seccion(prs, numero, titulo):
    s = _blank(prs)
    fondo = s.shapes.add_shape(1, 0, 0, ANCHO, ALTO)
    fondo.fill.solid()
    fondo.fill.fore_color.rgb = RGBColor(0xF2, 0xF5, 0xEE)
    fondo.line.fill.background()
    fondo.shadow.inherit = False
    _caja(s, Inches(1.2), Inches(2.9), Inches(1.2), Inches(1.4), str(numero),
          size=64, bold=True, color=VERDE_CLARO)
    _caja(s, Inches(2.6), Inches(3.1), Inches(9.6), Inches(1.2), titulo,
          size=36, bold=True, color=VERDE)
    return s


def slide_titulo(prs, titulo, bajada=""):
    """Slide de contenido con encabezado; devuelve (slide, y_contenido)."""
    s = _blank(prs)
    _caja(s, Inches(0.7), Inches(0.42), Inches(12.0), Inches(0.75), titulo,
          size=30, bold=True, color=VERDE)
    y = Inches(1.28)
    if bajada:
        _caja(s, Inches(0.7), Inches(1.18), Inches(12.0), Inches(0.5), bajada,
              size=15, color=GRIS_CLARO)
        y = Inches(1.78)
    linea = s.shapes.add_shape(1, Inches(0.7), Inches(1.12), Inches(1.5), Inches(0.035))
    linea.fill.solid()
    linea.fill.fore_color.rgb = VERDE_CLARO
    linea.line.fill.background()
    linea.shadow.inherit = False
    return s, y


def vinetas(slide, x, y, w, items, size=17, espacio=Pt(13)):
    """Lista con viñetas; cada item puede ser (texto, negrita) o solo texto."""
    tb = slide.shapes.add_textbox(x, y, w, Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        texto, bold = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = espacio
        punto = p.add_run()
        punto.text = "▪  "
        punto.font.size = Pt(size)
        punto.font.color.rgb = VERDE_CLARO
        punto.font.bold = True
        r = p.add_run()
        r.text = texto
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = GRIS
    return tb


def imagen(slide, path: Path, x, y, ancho_max=None, alto_max=None):
    """Inserta una imagen respetando su proporcion dentro de la caja dada."""
    from PIL import Image as PILImage

    w_px, h_px = PILImage.open(path).size
    prop = h_px / w_px
    if ancho_max and alto_max:
        w = ancho_max
        h = Emu(int(w * prop))
        if h > alto_max:
            h = alto_max
            w = Emu(int(h / prop))
    elif ancho_max:
        w, h = ancho_max, Emu(int(ancho_max * prop))
    else:
        h, w = alto_max, Emu(int(alto_max / prop))
    return slide.shapes.add_picture(str(path), x, y, width=w, height=h)


def imagen_centrada(slide, path: Path, y, ancho_max, alto_max):
    pic = imagen(slide, path, Inches(0), y, ancho_max, alto_max)
    pic.left = Emu(int((ANCHO - pic.width) / 2))
    return pic


def tabla(slide, x, y, w, h, datos: list[list[str]], anchos=None,
          size=13, destacar_fila=None, destacar_color=None):
    """Tabla con la primera fila como encabezado."""
    filas, cols = len(datos), len(datos[0])
    forma = slide.shapes.add_table(filas, cols, x, y, w, h)
    tbl = forma.table

    if anchos:
        total = sum(anchos)
        for j, a in enumerate(anchos):
            tbl.columns[j].width = Emu(int(w * a / total))

    for i, fila in enumerate(datos):
        tbl.rows[i].height = Inches(0.42 if i == 0 else 0.36)
        for j, celda in enumerate(fila):
            c = tbl.cell(i, j)
            c.text = str(celda)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = Inches(0.09)
            c.margin_right = Inches(0.09)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(size if i else size)
                r.font.name = "Calibri"
                r.font.bold = i == 0
                r.font.color.rgb = BLANCO if i == 0 else GRIS
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = VERDE
            elif destacar_fila is not None and i == destacar_fila:
                c.fill.fore_color.rgb = destacar_color or RGBColor(0xE8, 0xF1, 0xDC)
            else:
                c.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i % 2 else RGBColor(0xF6, 0xF7, 0xF4)
    return tbl


def tarjeta(slide, x, y, w, h, titulo, valor, detalle="", color=VERDE):
    """Tarjeta tipo KPI: un numero grande con etiqueta."""
    caja = slide.shapes.add_shape(1, x, y, w, h)
    caja.fill.solid()
    caja.fill.fore_color.rgb = RGBColor(0xF6, 0xF7, 0xF4)
    caja.line.color.rgb = RGBColor(0xDD, 0xE2, 0xD8)
    caja.line.width = Pt(1)
    caja.shadow.inherit = False
    caja.text_frame.text = ""

    borde = slide.shapes.add_shape(1, x, y, Inches(0.055), h)
    borde.fill.solid()
    borde.fill.fore_color.rgb = color
    borde.line.fill.background()
    borde.shadow.inherit = False

    _caja(slide, x + Inches(0.28), y + Inches(0.16), w - Inches(0.45), Inches(0.35),
          titulo, size=12, bold=True, color=GRIS_CLARO)
    _caja(slide, x + Inches(0.28), y + Inches(0.52), w - Inches(0.45), Inches(0.62),
          valor, size=30, bold=True, color=color)
    if detalle:
        _caja(slide, x + Inches(0.28), y + h - Inches(0.62), w - Inches(0.45), Inches(0.5),
              detalle, size=11, color=GRIS_CLARO)


def nota(slide, texto, y=None, color=GRIS_CLARO, size=13):
    """Linea de cierre al pie del slide: el 'take-away'."""
    y = y or Inches(6.72)
    barra = slide.shapes.add_shape(1, Inches(0.7), y, Inches(0.05), Inches(0.42))
    barra.fill.solid()
    barra.fill.fore_color.rgb = VERDE_CLARO
    barra.line.fill.background()
    barra.shadow.inherit = False
    _caja(slide, Inches(0.92), y - Inches(0.02), Inches(11.7), Inches(0.5), texto,
          size=size, color=color, anchor=MSO_ANCHOR.MIDDLE)


def numerar(prs) -> None:
    """Numera todos los slides salvo la portada."""
    total = len(prs.slides._sldIdLst)
    for i, s in enumerate(prs.slides, start=1):
        if i == 1:
            continue
        _caja(s, ANCHO - Inches(1.15), ALTO - Inches(0.55), Inches(0.7), Inches(0.35),
              f"{i}/{total}", size=11, color=GRIS_CLARO, align=PP_ALIGN.RIGHT)
